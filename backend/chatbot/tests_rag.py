"""Tests for the RAG layer (chatbot/rag/) and its orchestration (DocumentService).

Three levels:
    EmbedderTests        — pure Python, no database (SimpleTestCase);
    RAGServiceTests      — index()/retrieve() against the real pgvector table;
    DocumentServiceTests — File -> ProcessedChunks -> chunks + vectors, using
                           a real DOCX built in memory.

Run: python manage.py test chatbot --keepdb
"""
from __future__ import annotations

import io
import math
import uuid

import docx  # python-docx: builds a genuine .docx for the orchestration test
from django.db import connection
from django.test import SimpleTestCase, TestCase

from chatbot.models import ChatSession, Document, DocumentChunk, DocumentVector
from chatbot.rag.embedder import DEFAULT_DIMENSION, Embedder
from chatbot.rag.service import RAGService
from chatbot.rag.vector_store import VectorStore
from chatbot.services.document_processor import (
    DocumentProcessingError,
    ProcessedChunk,
)
from chatbot.services.document_service import DocumentService


def make_chunk(
    document: Document,
    index: int,
    text: str,
    *,
    page: int | None = None,
    heading: str | None = None,
    session: ChatSession | None = None,
) -> ProcessedChunk:
    """A ProcessedChunk shaped exactly like DocumentService produces them."""
    chunk = ProcessedChunk(
        chunk_id=str(uuid.uuid4()),
        document_id=str(document.id),
        chunk_index=index,
        text=text,
        metadata={
            "source_file": document.file_name,
            "file_type": "pdf",
            "page": page,
            "heading": heading,
            "char_count": len(text),
        },
    )
    # Document-level metadata added by the orchestration layer before index()
    chunk.metadata["document_type"] = document.source
    chunk.metadata["session_id"] = str(session.id) if session else None
    return chunk


class EmbedderTests(SimpleTestCase):
    """The embedder must be deterministic, normalized and Django-free."""

    def test_dimension_and_unit_norm(self) -> None:
        embedder = Embedder()
        self.assertEqual(embedder.dimension, DEFAULT_DIMENSION)
        vector = embedder.embed("Chapter 3 should describe the methodology.")
        self.assertEqual(len(vector), DEFAULT_DIMENSION)
        norm = math.sqrt(sum(value * value for value in vector))
        self.assertAlmostEqual(norm, 1.0, places=9)

    def test_deterministic_across_instances(self) -> None:
        # Different instances (and therefore different processes) must agree:
        # the indexed vectors and a later query must share one space.
        text = "Les résultats du chapitre 3 sont présentés ici."
        self.assertEqual(Embedder().embed(text), Embedder().embed(text))

    def test_frozen_vector_regression_guard(self) -> None:
        # Exact bucket/sign assignment for a fixed input. If this fails, the
        # hashing scheme changed: every stored vector needs a re-index.
        expected = [
            0.0, -0.3333333333333333, 0.0, 0.3333333333333333,
            -0.3333333333333333, -0.6666666666666666, -0.3333333333333333,
            0.0, 0.0, 0.0, 0.3333333333333333, 0.0, 0.0, 0.0, 0.0, 0.0,
        ]
        actual = Embedder(16).embed("Chapter 3 must contain the methodology")
        self.assertEqual(len(actual), 16)
        for got, want in zip(actual, expected):
            self.assertAlmostEqual(got, want, places=12)

    def test_empty_and_stopword_only_texts_embed_to_zero_vector(self) -> None:
        embedder = Embedder()
        self.assertFalse(any(embedder.embed("")))
        self.assertFalse(any(embedder.embed("the and of le la des")))

    def test_related_text_outscores_unrelated_text(self) -> None:
        embedder = Embedder()

        def cosine(a: list[float], b: list[float]) -> float:
            return sum(x * y for x, y in zip(a, b))

        query = embedder.embed("What should Chapter 3 contain?")
        related = embedder.embed("Chapter 3 presents the methodology.")
        unrelated = embedder.embed("The cafeteria menu changes every Monday.")
        self.assertGreater(cosine(query, related), cosine(query, unrelated))

    def test_rejects_silly_dimensions(self) -> None:
        with self.assertRaises(ValueError):
            Embedder(dimension=4)

    def test_embed_batch_matches_single_embeds(self) -> None:
        embedder = Embedder()
        texts = ["first chunk", "", "second chunk"]
        self.assertEqual(
            embedder.embed_batch(texts), [embedder.embed(text) for text in texts]
        )


class RAGServiceTests(TestCase):
    """index() writes vectors + metadata; retrieve() only ever reads."""

    def setUp(self) -> None:
        self.session = ChatSession.objects.create(title="rag tests")
        self.uni = Document.objects.create(
            title="Guidelines",
            file_name="guidelines.pdf",
            file_path="university/guidelines.pdf",
            source="university",
            session=self.session,
        )
        self.stu = Document.objects.create(
            title="Thesis",
            file_name="thesis.pdf",
            file_path="student/thesis.pdf",
            source="student",
            session=self.session,
        )
        self.rag = RAGService()

    def _sample_chunks(self) -> list[ProcessedChunk]:
        return [
            make_chunk(
                self.uni, 0,
                "Chapter 3 must contain the methodology and the sampling plan.",
                page=3, heading="Methodology", session=self.session,
            ),
            make_chunk(
                self.uni, 1,
                "All submissions use the APA reference style.",
                page=7, heading="References", session=self.session,
            ),
            make_chunk(
                self.stu, 0,
                "The thesis abstract summarizes the research problem.",
                page=1, heading="Abstract", session=self.session,
            ),
        ]

    # ------------------------------------------------------------------ #
    # index()                                                             #
    # ------------------------------------------------------------------ #

    def test_index_stores_vectors_with_required_metadata(self) -> None:
        chunks = self._sample_chunks()
        written = self.rag.index(chunks)

        self.assertEqual(written, 3)
        self.assertEqual(DocumentVector.objects.count(), 3)

        # values() deliberately omits the embedding column (raw SQL owns it)
        row = (
            DocumentVector.objects.filter(chunk_id=chunks[0].chunk_id)
            .values(
                "document_id", "document_type", "session_id", "page",
                "section", "content", "chunk_index", "metadata",
            )
            .get()
        )
        self.assertEqual(row["document_id"], self.uni.id)
        self.assertEqual(row["document_type"], "university")
        self.assertEqual(row["session_id"], self.session.id)
        self.assertEqual(row["page"], 3)
        self.assertEqual(row["section"], "Methodology")
        self.assertEqual(row["content"], chunks[0].text)
        self.assertEqual(row["chunk_index"], 0)
        self.assertEqual(row["metadata"]["source_file"], "guidelines.pdf")

        with connection.cursor() as cursor:
            cursor.execute("SELECT vector_dims(embedding) FROM document_vectors LIMIT 1")
            self.assertEqual(cursor.fetchone()[0], DEFAULT_DIMENSION)

    def test_index_replaces_previous_vectors_of_the_same_document(self) -> None:
        chunks = [c for c in self._sample_chunks() if c.document_id == str(self.uni.id)]
        self.rag.index(chunks)
        self.rag.index(chunks)  # re-index: no duplicates

        self.assertEqual(DocumentVector.objects.filter(document=self.uni).count(), 2)
        self.assertEqual(DocumentVector.objects.count(), 2)

    # ------------------------------------------------------------------ #
    # retrieve()                                                          #
    # ------------------------------------------------------------------ #

    def test_retrieve_returns_best_match_first_with_source_metadata(self) -> None:
        self.rag.index(self._sample_chunks())

        hits = self.rag.retrieve("What should Chapter 3 contain?", top_k=3)

        self.assertTrue(hits)
        best = hits[0]
        self.assertEqual(best.document_id, str(self.uni.id))
        self.assertEqual(best.document_type, "university")
        self.assertEqual(best.page, 3)
        self.assertEqual(best.section, "Methodology")
        self.assertEqual(best.session_id, str(self.session.id))
        self.assertEqual(best.metadata["source_file"], "guidelines.pdf")
        self.assertGreater(best.score, 0.0)

        scores = [hit.score for hit in hits]
        self.assertEqual(scores, sorted(scores, reverse=True))
        self.assertLessEqual(len(hits), 3)

    def test_retrieve_filters_by_document_type(self) -> None:
        self.rag.index(self._sample_chunks())

        hits = self.rag.retrieve(
            "research methodology",
            filters={"document_type": "student"},
            top_k=5,
        )

        self.assertTrue(hits)
        for hit in hits:
            self.assertEqual(hit.document_id, str(self.stu.id))
            self.assertEqual(hit.document_type, "student")

    def test_retrieve_filters_by_session(self) -> None:
        self.rag.index(self._sample_chunks())

        hits = self.rag.retrieve(
            "methodology sampling",
            filters={"session_id": str(self.session.id)},
        )
        self.assertTrue(hits)
        self.assertTrue(all(h.session_id == str(self.session.id) for h in hits))

        none_left = self.rag.retrieve(
            "methodology sampling",
            filters={"session_id": str(uuid.uuid4())},
        )
        self.assertEqual(none_left, [])

    def test_retrieve_never_indexes_the_question(self) -> None:
        self.rag.index(self._sample_chunks())
        before = DocumentVector.objects.count()

        self.rag.retrieve("Why is the response rate so low in this survey?")

        self.assertEqual(DocumentVector.objects.count(), before)
        self.assertFalse(
            DocumentVector.objects.filter(
                content__icontains="response rate so low"
            ).exists()
        )

    def test_retrieve_with_blank_or_zero_query_returns_empty(self) -> None:
        self.rag.index(self._sample_chunks())
        before = DocumentVector.objects.count()

        self.assertEqual(self.rag.retrieve(""), [])
        self.assertEqual(self.rag.retrieve("   "), [])
        self.assertEqual(self.rag.retrieve("the and of"), [])  # stopword-only

        self.assertEqual(DocumentVector.objects.count(), before)

    def test_retrieve_rejects_unknown_filter_and_bad_top_k(self) -> None:
        self.rag.index(self._sample_chunks())

        with self.assertRaises(ValueError):
            self.rag.retrieve("methodology", filters={"title": "nope"})
        with self.assertRaises(ValueError):
            self.rag.retrieve("methodology", top_k=0)

    def test_delete_for_document_removes_only_that_documents_vectors(self) -> None:
        self.rag.index(self._sample_chunks())

        removed = VectorStore().delete_for_document(str(self.uni.id))

        self.assertEqual(removed, 2)
        self.assertEqual(DocumentVector.objects.filter(document=self.uni).count(), 0)
        self.assertEqual(DocumentVector.objects.filter(document=self.stu).count(), 1)


class DocumentServiceTests(TestCase):
    """File -> ProcessedChunks -> document_chunks + vectors, statuses included."""

    def setUp(self) -> None:
        self.session = ChatSession.objects.create(title="service tests")

    @staticmethod
    def make_docx_bytes() -> bytes:
        """A genuine .docx (python-docx) with a heading and two paragraphs."""
        document = docx.Document()
        document.add_heading("Chapter 3: Methodology", level=1)
        document.add_paragraph(
            "Chapter 3 must contain the methodology, the sampling strategy "
            "and the response rate analysis of the survey."
        )
        document.add_paragraph(
            "The university library keeps every submitted thesis for ten years."
        )
        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()

    def test_process_document_persists_chunks_and_indexes_vectors(self) -> None:
        document = Document.objects.create(
            title="Thesis.docx",
            file_name="Thesis.docx",
            file_path="student/Thesis.docx",
            source="student",
            session=self.session,
        )

        chunks = DocumentService().process_document(
            document, self.make_docx_bytes()
        )

        document.refresh_from_db()
        self.assertEqual(document.status, Document.Status.READY)
        self.assertTrue(chunks)
        self.assertEqual(
            DocumentChunk.objects.filter(document=document).count(), len(chunks)
        )
        self.assertEqual(
            DocumentVector.objects.filter(document=document).count(), len(chunks)
        )

        vector_row = (
            DocumentVector.objects.filter(document=document)
            .values("document_type", "session_id")
            .first()
        )
        self.assertEqual(vector_row["document_type"], "student")
        self.assertEqual(vector_row["session_id"], self.session.id)

        # End-to-end: the indexed vectors answer a retrieval query.
        hits = RAGService().retrieve(
            "sampling strategy response rate",
            filters={"document_type": "student"},
        )
        self.assertTrue(hits)
        self.assertEqual(hits[0].document_id, str(document.id))

    def test_unsupported_file_marks_document_failed(self) -> None:
        document = Document.objects.create(
            title="Legacy.doc",
            file_name="Legacy.doc",
            file_path="student/Legacy.doc",
            source="student",
            session=self.session,
        )

        with self.assertRaises(DocumentProcessingError):
            DocumentService().process_document(document, b"binary garbage")

        document.refresh_from_db()
        self.assertEqual(document.status, Document.Status.FAILED)
        self.assertEqual(DocumentChunk.objects.filter(document=document).count(), 0)
        self.assertEqual(DocumentVector.objects.filter(document=document).count(), 0)
